"""Q2 midnight context features; no observation on the target day is used.

The attachment reader intentionally reads only the header and 31 January rows.
The feature builder is reusable on other daily matrices, but does not fit or
update any rule. Missing lag days remain missing; no future-filled warm start.
"""
from __future__ import annotations

from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

JANUARY = pd.date_range("2025-01-01", "2025-01-31", name="date")
STEP_HOURS = 1 / 6
SLOTS_PER_DAY = 144
EPOCH = pd.Timestamp("2025-01-01")
FEATURE_GROUPS = {
    "season": ("season_sin", "season_cos"),
    "weekend": ("weekend",),
    "pv_mean_3d": ("pv_mean_3d",),
    "recency": ("calendar_day",),
    "pv_lag_1d": ("pv_lag_1d",),
    "pv_std_3d": ("pv_std_3d",),
    "pv_trend_3d": ("pv_trend_3d",),
    "load_mean_3d": ("load_mean_3d",),
}
FEATURE_DESCRIPTIONS = {
    "season": "月份圆周位置（sin/cos 为同一个季节特征的两维编码）",
    "weekend": "周六或周日标志；不判断法定节假日及调休",
    "pv_mean_3d": "前三个完整日的光伏日总电量均值（kWh）",
    "recency": "自 2025-01-01 起的日历天数；差值表示时间距离",
    "pv_lag_1d": "前一日光伏日总电量（kWh）",
    "pv_std_3d": "前三日光伏日总电量总体标准差（kWh，ddof=0）",
    "pv_trend_3d": "前一日与前三日光伏日总电量之差的一半（kWh/day）",
    "load_mean_3d": "前三个完整日的负载日总电量均值（kWh）",
}


def load_january(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Read January only from attachment 2, preserving its 144 ordered slots."""
    book = load_workbook(path, read_only=True, data_only=True)
    frames = {}
    try:
        for name, word in (("load", "负载"), ("pv", "光伏")):
            sheets = [s for s in book if word in s.title]
            if len(sheets) != 1:
                raise ValueError(f"Expected one sheet containing {word!r}")
            # max_row=32: February and later observations are never requested.
            rows = list(sheets[0].iter_rows(min_row=1, max_row=32, values_only=True))
            if len(rows) != 32 or len(rows[0]) != SLOTS_PER_DAY + 1:
                raise ValueError("Expected a header and 31 January rows of 144 slots")
            dates = pd.DatetimeIndex([r[0] for r in rows[1:]], name="date")
            if not dates.equals(JANUARY):
                raise ValueError("Attachment must start with ordered 2025 January days")
            labels = [str(v) for v in rows[0][1:]]
            minutes = []
            for label in labels:
                hh, mm, *_ = label.replace("+1", "").split(":")
                minutes.append(60 * int(hh) + int(mm) + (1440 if "+1" in label else 0))
            if minutes != list(range(10, 1441, 10)):
                raise ValueError("Expected right-endpoint labels 00:10 through 00:00+1")
            values = np.asarray([r[1:] for r in rows[1:]], dtype=float)
            if not np.isfinite(values).all() or (values < 0).any():
                raise ValueError("Missing, nonfinite or negative January source values")
            frames[name] = pd.DataFrame(values, index=dates, columns=labels)
    finally:
        book.close()
    if not frames["load"].columns.equals(frames["pv"].columns):
        raise ValueError("Load and PV slots must match")
    return frames["load"], frames["pv"]


def build_daily_features(load: pd.DataFrame, pv: pd.DataFrame,
                         target_dates=None) -> pd.DataFrame:
    """Build raw contexts at midnight; target_dates may extend past observations.

    Calendar-day reindexing ensures 'previous three days' never means three
    arbitrarily spaced rows. Statistics use shift(1) before any rolling window.
    Historical contexts are computed as of their own midnight as well.
    """
    if not load.index.equals(pv.index) or not load.columns.equals(pv.columns):
        raise ValueError("Aligned load/PV daily matrices are required")
    if load.empty or not isinstance(load.index, pd.DatetimeIndex):
        raise ValueError("Nonempty matrices with DatetimeIndex are required")
    if load.index.has_duplicates or not load.index.is_monotonic_increasing:
        raise ValueError("Dates must be unique and sorted")
    if load.index.tz is not None or not load.index.equals(load.index.normalize()):
        raise ValueError("Use timezone-naive midnight dates")
    if load.shape[1] != SLOTS_PER_DAY:
        raise ValueError("Expected 144 power slots per day")
    wanted = load.index if target_dates is None else pd.DatetimeIndex(target_dates)
    dates = pd.date_range(min(load.index.min(), wanted.min()),
                         max(load.index.max(), wanted.max()), name="date")
    daily_pv = (pv.sum(axis=1, min_count=SLOTS_PER_DAY) * STEP_HOURS).reindex(dates)
    daily_load = (load.sum(axis=1, min_count=SLOTS_PER_DAY) * STEP_HOURS).reindex(dates)
    lag_pv = daily_pv.shift(1)
    lag_load = daily_load.shift(1)
    angle = 2 * np.pi * (dates.month.to_numpy() - 1) / 12
    result = pd.DataFrame({
        "season_sin": np.sin(angle),
        "season_cos": np.cos(angle),
        "weekend": (dates.dayofweek >= 5).astype(float),
        "calendar_day": (dates - EPOCH).days.astype(float),
        "pv_mean_3d": lag_pv.rolling(3, min_periods=3).mean(),
        "pv_lag_1d": lag_pv,
        "pv_std_3d": lag_pv.rolling(3, min_periods=3).std(ddof=0),
        "pv_trend_3d": (daily_pv.shift(1) - daily_pv.shift(3)) / 2,
        "load_mean_3d": lag_load.rolling(3, min_periods=3).mean(),
    }, index=dates)
    return result.loc[wanted]


def features_from_attachment(path: str | Path) -> pd.DataFrame:
    """Convenience interface: attachment 2 to January raw midnight contexts."""
    return build_daily_features(*load_january(path))
