# -*- coding: utf-8 -*-
"""问题二全年情景赋权核：处境特征、冻结标准化、高斯权重与能量分数。

与 `solving/task4kernel/kernel.py` 接口兼容：`FrozenStandardizer`、`gaussian_weights`、
`energy_score`、`pairwise_distances`、`effective_scenarios`、`weights_for_day` 的
签名与语义保持一致，可直接替换调用。相对一月版新增：

1. `build_context_features` 给出全年处境的因果构造，含周同期（7 日滞后）与星期几结构；
2. `causal_pool` 实现"同季节且严格早于决策日"的候选池，季节由月份环形距离定义；
3. `KernelRule` 把特征集合、带宽、池规则与标尺策略打包成可序列化的冻结规则。

全部处境特征只使用决策日之前已完整结束的日数据，不含任何当日观测。
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

STEP_HOURS = 1.0 / 6.0
SLOTS_PER_DAY = 144
EPOCH = pd.Timestamp("2025-01-01")
WARMUP_DAYS = 7
SEASON_MONTH_RADIUS = 1
DATA_START = pd.Timestamp("2025-01-01")
DATA_END = pd.Timestamp("2025-12-31")

FEATURE_GROUPS: dict[str, tuple[str, ...]] = {
    "season": ("season_sin", "season_cos"),
    "weekday7": tuple(f"wd{i}" for i in range(7)),
    "weekend": ("weekend",),
    "low_load_day": ("low_load_day",),
    "net_lag_7d": ("net_lag_7d",),
    "load_lag_7d": ("load_lag_7d",),
    "pv_lag_7d": ("pv_lag_7d",),
    "load_mean_3d": ("load_mean_3d",),
    "pv_mean_3d": ("pv_mean_3d",),
    "net_mean_3d": ("net_mean_3d",),
    "net_mean_7d": ("net_mean_7d",),
    "load_std_7d": ("load_std_7d",),
    "pv_std_7d": ("pv_std_7d",),
}
FEATURE_DESCRIPTIONS = {
    "season": "月份圆周位置（sin/cos 是同一个季节特征的两维编码，取自然尺度）",
    "weekday7": "星期几的七维指示编码（同一星期几距离为 0，取自然尺度）",
    "weekend": "周六或周日标志；不判断法定节假日及调休",
    "low_load_day": "低负载日标志：按全年数据，低负载落在周五与周六（取自然尺度）",
    "net_lag_7d": "七个完整日前的净负荷日总电量（kWh），即上周同一星期几",
    "load_lag_7d": "七个完整日前的负载日总电量（kWh）",
    "pv_lag_7d": "七个完整日前的光伏日总电量（kWh）",
    "load_mean_3d": "前三个完整日的负载日总电量均值（kWh）",
    "pv_mean_3d": "前三个完整日的光伏日总电量均值（kWh）",
    "net_mean_3d": "前三个完整日的净负荷日总电量均值（kWh）",
    "net_mean_7d": "前七个完整日的净负荷日总电量均值（kWh）",
    "load_std_7d": "前七个完整日负载日总电量的总体标准差（kWh，ddof=0）",
    "pv_std_7d": "前七个完整日光伏日总电量的总体标准差（kWh，ddof=0）",
}
NATURAL_SCALE_COLUMNS = ("season_sin", "season_cos", "weekend", "low_load_day",
                         *FEATURE_GROUPS["weekday7"])


def columns_for(groups) -> list[str]:
    """把语义特征组展开为核距离实际使用的数值坐标。"""
    return [column for group in groups for column in FEATURE_GROUPS[group]]


def load_attachment(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """读取附件2 全年负载与光伏功率（kW），返回两个 365×144 的日矩阵。"""
    book = load_workbook(path, read_only=True, data_only=True)
    frames: dict[str, pd.DataFrame] = {}
    try:
        for name, word in (("load", "负载"), ("pv", "光伏")):
            sheets = [s for s in book if word in s.title]
            if len(sheets) != 1:
                raise ValueError(f"附件2 中应恰有一张含 {word!r} 的工作表")
            rows = list(sheets[0].iter_rows(min_row=1, max_row=366, values_only=True))
            if len(rows) != 366 or len(rows[0]) != SLOTS_PER_DAY + 1:
                raise ValueError("附件2 应为表头加 365 天、每天 144 个时段")
            dates = pd.DatetimeIndex([r[0] for r in rows[1:]], name="date")
            if not dates.equals(pd.date_range(DATA_START, DATA_END, name="date")):
                raise ValueError("附件2 日期列应为连续完整的 2025 全年")
            labels = [str(v) for v in rows[0][1:]]
            values = np.asarray([r[1:] for r in rows[1:]], dtype=float)
            if not np.isfinite(values).all() or (values < 0).any():
                raise ValueError("附件2 存在缺失、非有限或负功率")
            frames[name] = pd.DataFrame(values, index=dates, columns=labels)
    finally:
        book.close()
    if not frames["load"].columns.equals(frames["pv"].columns):
        raise ValueError("负载与光伏的时段标签必须一致")
    return frames["load"], frames["pv"]


def build_context_features(load_power: pd.DataFrame, pv_power: pd.DataFrame,
                           target_dates=None) -> pd.DataFrame:
    """构造每天 0:00 可知的处境特征；target_dates 可超出已观测范围。

    load_power、pv_power 为逐时段功率（kW）。所有统计量先 shift(1) 再做滚动，
    因此第 d 天的处境绝不包含第 d 天及之后的任何观测。日历日重索引确保
    "前三个完整日"沿自然日推进，而不是沿任意间隔的行推进。
    """
    if not load_power.index.equals(pv_power.index) or not load_power.columns.equals(pv_power.columns):
        raise ValueError("负载与光伏日矩阵必须对齐")
    if load_power.empty or not isinstance(load_power.index, pd.DatetimeIndex):
        raise ValueError("需要非空且带 DatetimeIndex 的日矩阵")
    if load_power.index.has_duplicates or not load_power.index.is_monotonic_increasing:
        raise ValueError("日期必须唯一且升序")
    if load_power.shape[1] != SLOTS_PER_DAY:
        raise ValueError("每天应为 144 个时段")
    wanted = load_power.index if target_dates is None else pd.DatetimeIndex(target_dates)
    dates = pd.date_range(min(load_power.index.min(), wanted.min()),
                          max(load_power.index.max(), wanted.max()), name="date")
    daily_load = (load_power.sum(axis=1, min_count=SLOTS_PER_DAY) * STEP_HOURS).reindex(dates)
    daily_pv = (pv_power.sum(axis=1, min_count=SLOTS_PER_DAY) * STEP_HOURS).reindex(dates)
    daily_net = daily_load - daily_pv
    angle = 2 * np.pi * (dates.month.to_numpy() - 1) / 12
    weekday = dates.dayofweek.to_numpy()
    result = pd.DataFrame({
        "season_sin": np.sin(angle),
        "season_cos": np.cos(angle),
        "weekend": (weekday >= 5).astype(float),
        "low_load_day": ((weekday >= 4) & (weekday <= 5)).astype(float),
        **{f"wd{i}": (weekday == i).astype(float) for i in range(7)},
        "net_lag_7d": daily_net.shift(7),
        "load_lag_7d": daily_load.shift(7),
        "pv_lag_7d": daily_pv.shift(7),
        "load_mean_3d": daily_load.shift(1).rolling(3, min_periods=3).mean(),
        "pv_mean_3d": daily_pv.shift(1).rolling(3, min_periods=3).mean(),
        "net_mean_3d": daily_net.shift(1).rolling(3, min_periods=3).mean(),
        "net_mean_7d": daily_net.shift(1).rolling(7, min_periods=7).mean(),
        "load_std_7d": daily_load.shift(1).rolling(7, min_periods=7).std(ddof=0),
        "pv_std_7d": daily_pv.shift(1).rolling(7, min_periods=7).std(ddof=0),
    }, index=dates)
    return result.loc[wanted]


def month_ring_distance(month_a: np.ndarray, month_b: np.ndarray) -> np.ndarray:
    """月份在 12 个月环上的最短距离，用于定义同季节。"""
    gap = np.abs(np.asarray(month_a) - np.asarray(month_b))
    return np.minimum(gap, 12 - gap)


def causal_pool(dates: pd.DatetimeIndex, day: pd.Timestamp, features: pd.DataFrame,
                columns, radius: int = SEASON_MONTH_RADIUS,
                warmup: int = WARMUP_DAYS) -> pd.DatetimeIndex:
    """返回第 day 天的候选历史池：严格早于 day、同季节、且处境特征完整。"""
    dates, day = pd.DatetimeIndex(dates), pd.Timestamp(day)
    if day not in dates:
        raise ValueError("决策日必须落在数据日期索引内")
    index = dates.get_indexer([day])[0]
    past = dates[:index]
    if len(past) == 0:
        return pd.DatetimeIndex([], name="date")
    keep = month_ring_distance(past.month.to_numpy(), np.full(len(past), day.month)) <= radius
    keep &= np.arange(len(past)) >= warmup
    pool = past[keep]
    if len(pool):
        usable = features.loc[pool, list(columns)].notna().all(axis=1).to_numpy()
        pool = pool[usable]
    return pool


@dataclass(frozen=True)
class FrozenStandardizer:
    """参考池总体均值/标准差，冻结后不再随查询日更新（与 task4kernel 一致）。"""

    columns: tuple[str, ...]
    mean: tuple[float, ...]
    scale: tuple[float, ...]

    @classmethod
    def fit(cls, reference: pd.DataFrame, fixed_scale_columns=()) -> "FrozenStandardizer":
        values = reference.to_numpy(dtype=float)
        if len(values) < 2 or not np.isfinite(values).all():
            raise ValueError("至少需要两行完整的有限参考处境")
        mean, scale = values.mean(axis=0), values.std(axis=0, ddof=0)
        for i, column in enumerate(reference.columns):
            if column in fixed_scale_columns:
                mean[i], scale[i] = 0.0, 1.0
        scale = np.where(scale > 1e-12, scale, 1.0)
        return cls(tuple(reference.columns), tuple(mean), tuple(scale))

    def transform(self, frame: pd.DataFrame) -> pd.DataFrame:
        values = frame.loc[:, list(self.columns)].to_numpy(dtype=float)
        if not np.isfinite(values).all():
            raise ValueError("处境不完整：缺少因果滞后观测")
        return pd.DataFrame((values - self.mean) / self.scale,
                            index=frame.index, columns=self.columns)

    def to_dict(self) -> dict:
        return {"columns": list(self.columns), "mean": list(self.mean),
                "scale": list(self.scale)}

    @classmethod
    def from_dict(cls, data: dict) -> "FrozenStandardizer":
        return cls(tuple(data["columns"]), tuple(data["mean"]), tuple(data["scale"]))


def gaussian_weights(history_features, target_features, h: float) -> np.ndarray:
    """对同一候选池的一个或多个目标返回归一化高斯权重（与 task4kernel 一致）。

    history_features: (K, p)；target_features: (p,) 或 (D, p)。
    h=inf 显式表示等权极限；减去最小平方距离以避免指数下溢。
    """
    history = np.asarray(history_features, dtype=float)
    target = np.asarray(target_features, dtype=float)
    single = target.ndim == 1
    if history.ndim != 2 or target.ndim not in (1, 2) or not len(history):
        raise ValueError("需要非空 (K,p) 历史与 (p,) 或 (D,p) 目标")
    target = np.atleast_2d(target)
    if history.shape[1] != target.shape[1]:
        raise ValueError("特征维度必须一致")
    if not np.isfinite(history).all() or not np.isfinite(target).all():
        raise ValueError("特征必须为有限值")
    if np.isnan(h) or h <= 0:
        raise ValueError("带宽必须为正（或正无穷）")
    if np.isinf(h):
        weights = np.full((len(target), len(history)), 1 / len(history))
    else:
        squared = ((target[:, None, :] - history[None, :, :]) ** 2).sum(axis=2)
        shifted = squared - squared.min(axis=1, keepdims=True)
        with np.errstate(over="ignore"):
            logits = -0.5 * (shifted / h) / h
        unnormalized = np.exp(logits)
        weights = unnormalized / unnormalized.sum(axis=1, keepdims=True)
    return weights[0] if single else weights


def effective_scenarios(weights) -> np.ndarray:
    """K_eff=1/sum(pi^2)，只作集中度诊断，不参与调参。"""
    weights = np.asarray(weights, dtype=float)
    return 1 / np.sum(weights ** 2, axis=-1)


def pairwise_distances(x, y=None) -> np.ndarray:
    """精确欧氏距离；不做坐标标准化，也不做均值预测。"""
    x = np.asarray(x, dtype=float)
    y = x if y is None else np.asarray(y, dtype=float)
    return np.linalg.norm(x[:, None, :] - y[None, :, :], axis=-1)


def energy_score(scenarios, truth, weights) -> float:
    """完整离散能量分数，含负的二分之一成对距离项。"""
    x, y, w = (np.asarray(a, dtype=float) for a in (scenarios, truth, weights))
    if x.ndim != 2 or y.shape != (x.shape[1],) or w.shape != (len(x),):
        raise ValueError("需要 (K,T) 情景、(T,) 真值与 (K,) 权重")
    if not all(np.isfinite(v).all() for v in (x, y, w)):
        raise ValueError("ES 输入必须为有限值")
    if (w < 0).any() or not np.isclose(w.sum(), 1, rtol=0, atol=1e-12):
        raise ValueError("ES 需要归一化的非负概率")
    return float(w @ np.linalg.norm(x - y, axis=1) - 0.5 * w @ pairwise_distances(x) @ w)


def weights_for_day(feature_frame: pd.DataFrame, target_day, history_days,
                    scaler: FrozenStandardizer, h: float) -> pd.Series:
    """日期安全的下游接口；候选池由调用者显式给出（签名与 task4kernel 一致）。"""
    day, dates = pd.Timestamp(target_day), pd.DatetimeIndex(history_days)
    if dates.empty or dates.has_duplicates or not (dates < day).all():
        raise ValueError("每个唯一情景日期都必须早于决策日")
    historical = scaler.transform(feature_frame.loc[dates])
    target = scaler.transform(feature_frame.loc[[day]])
    return pd.Series(gaussian_weights(historical, target.iloc[0], h),
                     index=dates, name="weight")


@dataclass(frozen=True)
class KernelRule:
    """冻结后的赋权规则：特征集合、带宽、池参数与标尺拟合策略。"""

    groups: tuple[str, ...]
    bandwidth: float
    season_month_radius: int = SEASON_MONTH_RADIUS
    warmup_days: int = WARMUP_DAYS
    scaling: str = "rolling_pool"
    adopted: bool = True

    def columns(self) -> list[str]:
        return columns_for(self.groups)

    def pool(self, dates, day, features) -> pd.DatetimeIndex:
        return causal_pool(dates, day, features, self.columns(),
                           self.season_month_radius, self.warmup_days)

    def weights(self, features: pd.DataFrame, day, dates=None) -> pd.Series:
        """按规则给出第 day 天在候选池上的权重；标尺只用池内处境拟合。"""
        day = pd.Timestamp(day)
        pool = self.pool(features.index if dates is None else dates, day, features)
        if pool.empty:
            raise ValueError(f"{day.date()} 的候选池为空")
        columns = self.columns()
        scaler = FrozenStandardizer.fit(features.loc[pool, columns], NATURAL_SCALE_COLUMNS)
        return weights_for_day(features, day, pool, scaler, self.bandwidth)

    def to_dict(self) -> dict:
        return {"schema_version": 1, "groups": list(self.groups),
                "bandwidth": None if np.isinf(self.bandwidth) else float(self.bandwidth),
                "bandwidth_limit": "uniform" if np.isinf(self.bandwidth) else "finite",
                "season_month_radius": self.season_month_radius,
                "warmup_days": self.warmup_days, "scaling": self.scaling,
                "adopted": self.adopted,
                "natural_scale_columns": list(NATURAL_SCALE_COLUMNS)}

    @classmethod
    def from_dict(cls, data: dict) -> "KernelRule":
        h = data["bandwidth"]
        return cls(tuple(data["groups"]), float("inf") if h is None else float(h),
                   data["season_month_radius"], data["warmup_days"],
                   data["scaling"], data.get("adopted", True))
