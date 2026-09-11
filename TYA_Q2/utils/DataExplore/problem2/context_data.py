"""第二问理解对齐：读取原始附件，核对时序、能量规模及因果预测基线。"""
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

DELTA_H = 1 / 6
TEST_START = pd.Timestamp('2025-02-01')
SELECTED_DAYS = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']


def label_minutes(labels):
    """统一 Excel time 与字符串时间为分钟序号，仅规范格式，不平移时刻。"""
    result = []
    for label in labels:
        value = str(label)
        pieces = value.replace('+1', '').split(':')
        result.append(int(pieces[0]) * 60 + int(pieces[1]) + (1440 if '+1' in value else 0))
    return result


def inspect_sources(root):
    """返回附件1、2和结果模板各表尺寸、头部与尾部，保留原始时间标签。"""
    result = {}
    for relative in ['附件1.xlsx', '附件2.xlsx', '附件5/result2.xlsx']:
        book = load_workbook(Path(root) / 'ProblemC/附件' / relative, read_only=True, data_only=True)
        result[relative] = {}
        for sheet in book:
            rows = list(sheet.values)
            result[relative][sheet.title] = {
                'shape': [sheet.max_row, sheet.max_column],
                'top_left': [list(row[:5]) for row in rows[:4]],
                'top_right': [list(row[-3:]) for row in rows[:2]],
                'bottom_left': [list(row[:5]) for row in rows[-2:]],
            }
        book.close()
    return result


def load_inputs(root):
    """读取附件1与附件2；按原始记录顺序组织，暂不裁定时段起止标签。"""
    folder = Path(root) / 'ProblemC/附件'
    typical = pd.read_excel(folder / '附件1.xlsx')
    actual = pd.read_excel(folder / '附件2.xlsx', sheet_name=None, index_col=0)
    return typical, actual


def audit_matrix(frame):
    """返回矩阵完整性、日期连续性、非负性和标签端点。"""
    values = frame.to_numpy(dtype=float)
    dates = pd.to_datetime(frame.index)
    return {'shape': list(frame.shape), 'missing': int(np.isnan(values).sum()),
            'negative': int((values < 0).sum()), 'nonfinite': int((~np.isfinite(values)).sum()),
            'duplicate_dates': int(dates.duplicated().sum()),
            'continuous_dates': bool(dates.equals(pd.date_range('2025-01-01', '2025-12-31'))),
            'first_label': str(frame.columns[0]), 'last_label': str(frame.columns[-1]),
            'min_kW': float(np.nanmin(values)), 'max_kW': float(np.nanmax(values))}


def energy_summary(load, pv):
    """以十分钟区间代表功率换算电量，量化净负荷规模及光伏盈余。"""
    net = load - pv
    daily = pd.DataFrame({'负载_kWh': load.sum(axis=1) * DELTA_H,
                          '光伏_kWh': pv.sum(axis=1) * DELTA_H,
                          '净用电_kWh': net.sum(axis=1) * DELTA_H,
                          '盈余时段数': (net < 0).sum(axis=1)})
    daily.index = pd.to_datetime(daily.index)
    stats = {'平均日负载_kWh': float(daily['负载_kWh'].mean()),
             '平均日光伏_kWh': float(daily['光伏_kWh'].mean()),
             '平均日净用电_kWh': float(daily['净用电_kWh'].mean()),
             '光伏盈余时段数': int((net < 0).to_numpy().sum()),
             '光伏盈余时段占比': float((net < 0).to_numpy().mean()),
             '最大盈余功率_kW': float((-net).to_numpy().max()),
             '盈余超过5000kW的时段数': int((net < -5000).to_numpy().sum())}
    return stats, daily.loc[pd.to_datetime(SELECTED_DAYS)]


def causal_baselines(load, pv):
    """仅用日前历史生成三种简单预测，比较2—12月WAPE和净负荷误差相关性。"""
    rows = []
    for name, frame in [('负载', load), ('光伏', pv), ('净负荷', load - pv)]:
        mask = pd.to_datetime(frame.index) >= TEST_START
        truth = frame.loc[mask].to_numpy(dtype=float)
        predictions = {'前一日同刻': frame.shift(1), '前一周同刻': frame.shift(7),
                       '过去7日同刻均值': frame.shift(1).rolling(7).mean()}
        for method, pred in predictions.items():
            error = pred.loc[mask].to_numpy(dtype=float) - truth
            rows.append({'对象': name, '方法': method, 'MAE_kW': np.abs(error).mean(),
                         'WAPE_pct': 100 * np.abs(error).sum() / np.abs(truth).sum()})
    pred_net = load.shift(7) - pv.shift(1).rolling(7).mean()
    mask = pd.to_datetime(load.index) >= TEST_START
    residual = ((load - pv) - pred_net).loc[mask].to_numpy(dtype=float)
    error_info = {'组合基线净负荷_MAE_kW': float(np.abs(residual).mean()),
                  '组合基线净负荷_WAPE_pct': float(100 * np.abs(residual).sum() / np.abs((load-pv).loc[mask].to_numpy()).sum()),
                  '日内相邻误差相关系数': float(np.corrcoef(residual[:, :-1].ravel(), residual[:, 1:].ravel())[0, 1]),
                  '测试天数': int(mask.sum())}
    return pd.DataFrame(rows), error_info
