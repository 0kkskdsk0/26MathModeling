"""结果导出：按附件 5 模板写出 result3.xlsx。"""
import numpy as np
import pandas as pd
import openpyxl

import myh.src.config as config

BLOCK_LABELS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def _dates():
    return pd.date_range("2025-02-01", "2025-12-31", freq="D")


def _minutes_to_hhmm(m):
    return f"{int(m // 60):02d}:{int(m % 60):02d}"


def _emergency_records(E):
    records = []
    dates = _dates()
    for di in range(E.shape[0]):
        idx = np.flatnonzero(E[di] > 1e-9)
        if idx.size == 0:
            continue
        groups = np.split(idx, np.flatnonzero(np.diff(idx) != 1) + 1)
        for g in groups:
            start = int(g[0]) * 10
            end = int(g[-1] + 1) * 10
            records.append((dates[di],
                            f"{_minutes_to_hhmm(start)}-{_minutes_to_hhmm(end)}",
                            round(float(E[di, g].sum()), 6)))
    return records


def _fill_grid(ws, mat, P, G0=None):
    n = mat.shape[0]
    for di in range(n):
        row = di + 2
        for c in range(config.T):
            ws.cell(row=row, column=c + 2, value=round(float(mat[di, c]), 6))
        ws.cell(row=row, column=config.T + 2, value=round(float(mat[di].sum()), 6))
        cost = np.sum(P[di] * mat[di])
        if G0 is not None:
            cost += config.ADJUST_FRAC * np.sum(P[di] * np.abs(mat[di] - G0[di]))
        ws.cell(row=row, column=config.T + 3, value=round(float(cost), 6))


def write_result(template_path, output_path, ctrl, data, price_type="fixed", has_adjust=True):
    out_slice = slice(config.OUTPUT_START_IDX, config.OUTPUT_END_IDX + 1)
    if price_type == "varying":
        P = data.price_varying[out_slice]
    else:
        P = np.broadcast_to(data.price_fixed, (config.N_DAYS, config.T))[out_slice]
    G0 = ctrl.G0[out_slice]
    GF = ctrl.GF[out_slice]
    C = ctrl.C[out_slice]
    D = ctrl.D[out_slice]
    E = ctrl.E[out_slice]
    S = ctrl.S[out_slice]

    wb = openpyxl.load_workbook(template_path)
    _fill_grid(wb["计划购电量"], G0, P)
    if has_adjust:
        _fill_grid(wb["调整购电量"], GF, P, G0=G0)

    ws_cd = wb["充放电量"]
    ws_cd.delete_rows(2, ws_cd.max_row)
    dates = _dates()
    r = 2
    for di in range(C.shape[0]):
        for b in range(6):
            c0, c1 = b * 24, (b + 1) * 24
            ws_cd.cell(row=r, column=1, value=dates[di])
            ws_cd.cell(row=r, column=2, value=BLOCK_LABELS[b])
            ws_cd.cell(row=r, column=3, value=round(float(C[di, c0:c1].sum()), 6))
            ws_cd.cell(row=r, column=4, value=round(float(D[di, c0:c1].sum()), 6))
            if b == 0:
                ws_cd.cell(row=r, column=5, value="0:00")
                ws_cd.cell(row=r, column=6, value=round(float(S[di, 0]), 6))
            elif b == 1:
                ws_cd.cell(row=r, column=5, value="24:00")
                ws_cd.cell(row=r, column=6, value=round(float(S[di, config.T]), 6))
            r += 1

    ws_emg = wb["紧急购电量"]
    ws_emg.delete_rows(2, ws_emg.max_row)
    for i, (dt, seg, amount) in enumerate(_emergency_records(E)):
        ws_emg.cell(row=i + 2, column=1, value=dt.to_pydatetime())
        ws_emg.cell(row=i + 2, column=2, value=seg)
        ws_emg.cell(row=i + 2, column=3, value=amount)

    wb.save(output_path)
    return output_path
